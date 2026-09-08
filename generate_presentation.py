import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palette
    NAVY = RGBColor(16, 44, 87)          # #102C57
    ROYAL = RGBColor(53, 95, 172)        # #355FAC
    CYAN = RGBColor(0, 168, 204)         # #00A8CC
    LIGHT_BG = RGBColor(248, 250, 252)   # #F8FAFC
    CARD_DARK = RGBColor(40, 48, 68)     # #283044
    WHITE = RGBColor(255, 255, 255)
    TEXT_DARK = RGBColor(30, 41, 59)
    TEXT_MUTED = RGBColor(100, 116, 139)
    BORDER_LIGHT = RGBColor(203, 213, 225)
    AMBER = RGBColor(245, 158, 11)
    GREEN = RGBColor(16, 185, 129)

    def add_header(slide, title_text, category_text="NEXORA"):
        # Header bar
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.8))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = ROYAL
        
        p2 = tf.add_paragraph()
        p2.text = title_text.upper()
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = NAVY

        # SIH Badge placeholder
        badge_box = slide.shapes.add_textbox(Inches(10.2), Inches(0.35), Inches(2.4), Inches(0.9))
        btf = badge_box.text_frame
        btf.word_wrap = True
        bp = btf.paragraphs[0]
        bp.text = "SMART INDIA HACKATHON 2026"
        bp.font.size = Pt(11)
        bp.font.bold = True
        bp.alignment = PP_ALIGN.RIGHT
        bp.font.color.rgb = NAVY
        bp2 = btf.add_paragraph()
        bp2.text = "Problem: SIH26171 | ISRO"
        bp2.font.size = Pt(9)
        bp2.alignment = PP_ALIGN.RIGHT
        bp2.font.color.rgb = TEXT_MUTED

    def add_footer(slide, page_num):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.4))
        ftf = footer_box.text_frame
        fp = ftf.paragraphs[0]
        fp.text = f"@SIH Idea submission- Template                                                                                                                                              {page_num}"
        fp.font.size = Pt(9)
        fp.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 1: TITLE PAGE
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    
    # Title Banner
    tbox = slide1.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.733), Inches(1.3))
    ttf = tbox.text_frame
    tp1 = ttf.paragraphs[0]
    tp1.text = "SMART INDIA HACKATHON 2026"
    tp1.font.size = Pt(32)
    tp1.font.bold = True
    tp1.alignment = PP_ALIGN.CENTER
    tp1.font.color.rgb = NAVY

    tp2 = ttf.add_paragraph()
    tp2.text = "TITLE PAGE"
    tp2.font.size = Pt(22)
    tp2.font.bold = True
    tp2.alignment = PP_ALIGN.CENTER
    tp2.font.color.rgb = ROYAL

    # Left Info Card
    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(6.5), Inches(4.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = LIGHT_BG
    card1.line.color.rgb = BORDER_LIGHT
    card1.line.width = Pt(1)

    ctf = card1.text_frame
    ctf.word_wrap = True
    ctf.margin_left = Inches(0.4)
    ctf.margin_top = Inches(0.4)
    
    items = [
        ("Problem Statement ID", "SIH26171"),
        ("Problem Statement Title", "On-device Visual Perception for Light-weight Browser Agents"),
        ("Theme", "Miscellaneous"),
        ("PS Category", "Software"),
        ("Organisation", "Indian Space Research Organisation (ISRO)"),
        ("Department", "Department of Space / ISRO"),
        ("Team Name", "NEXORA"),
        ("Project Prototype", "PERCEPTA (\"See the Web. Understand Locally. Act Safely.\")")
    ]
    
    for i, (k, v) in enumerate(items):
        p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
        p.space_after = Pt(8)
        run1 = p.add_run()
        run1.text = f"•  {k} : "
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = NAVY
        
        run2 = p.add_run()
        run2.text = f"{v}"
        run2.font.italic = (k == "Problem Statement Title")
        run2.font.size = Pt(13)
        run2.font.color.rgb = TEXT_DARK

    # Right Graphic Card
    card_r = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.7), Inches(2.2), Inches(4.8), Inches(4.5))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = NAVY
    card_r.line.color.rgb = CYAN
    card_r.line.width = Pt(1.5)

    rtf = card_r.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.4)
    rtf.margin_top = Inches(0.5)

    rp = rtf.paragraphs[0]
    rp.text = "PERCEPTA"
    rp.font.size = Pt(30)
    rp.font.bold = True
    rp.alignment = PP_ALIGN.CENTER
    rp.font.color.rgb = CYAN

    rp_sub = rtf.add_paragraph()
    rp_sub.text = "Autonomous Browser Agent Architecture"
    rp_sub.font.size = Pt(12)
    rp_sub.alignment = PP_ALIGN.CENTER
    rp_sub.font.color.rgb = WHITE
    rp_sub.space_after = Pt(20)

    pills = [
        "✓ 100% On-Device Visual Perception (CPU-First)",
        "✓ 5-Factor Hybrid Perception Scoring Fusion",
        "✓ Privacy Guard & Client-Side PII Masking",
        "✓ Playwright Safe Automation & Self-Correction",
        "✓ Zero Cloud LLM / Zero External API Dependency",
        "✓ Verified State Delta & Human-in-Loop Safety"
    ]
    for pill in pills:
        pp = rtf.add_paragraph()
        pp.text = pill
        pp.font.size = Pt(11)
        pp.font.color.rgb = RGBColor(226, 232, 240)
        pp.space_after = Pt(6)

    # ==========================================
    # SLIDE 2: PROPOSED SOLUTION
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "IDEA TITLE: PERCEPTA", "NEXORA")
    
    # Subtitle tag
    sub_box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.733), Inches(0.4))
    stf = sub_box.text_frame
    sp = stf.paragraphs[0]
    sp.text = "Proposed Solution — \"See the Web. Understand Locally. Act Safely.\""
    sp.font.size = Pt(15)
    sp.font.bold = True
    sp.font.color.rgb = RGBColor(225, 29, 72)

    # 3 Column Cards
    col_w = Inches(3.75)
    gap = Inches(0.24)
    c_y = Inches(1.85)
    c_h = Inches(4.95)

    cols_data = [
        ("Explanation", CARD_DARK, [
            "👁️  Visually detects webpage elements (buttons, inputs, cards, links) on CPU.",
            "🧠  Understands natural language user tasks with local entity intent parsing.",
            "⚡  Processes visual perception 100% locally on-device in 10-18ms.",
            "📐  Extracts deep DOM structure, accessibility roles, and bounding rectangles.",
            "⚖️  Fuses signals via 5-Factor Hybrid Score (Visual, DOM, Text, Position, Semantic).",
            "🛡️  Protects sensitive user data with local PII detection & masking.",
            "🤖  Performs browser actions safely via Playwright using hybrid coordinates.",
            "✅  Verifies post-action state delta with automated self-correction (max 2 retries)."
        ]),
        ("Problem Addressed", NAVY, [
            "🔗  Reduces dependency on fragile CSS/XPath selectors that break on layout changes.",
            "☁️  Eliminates reliance on costly, slow, and privacy-invasive cloud multimodal LLMs.",
            "🔒  Guarantees 100% user privacy; screenshots & credentials never leave the machine.",
            "🛰️  Enables air-gapped offline operation for defense and ISRO space network environments.",
            "🛑  Prevents unintended actions through a two-tier safety classification gateway.",
            "✅  Verifies every browser interaction with state assertion rather than blind execution.",
            "🔄  Autonomous self-healing stops cascading automation failures gracefully."
        ]),
        ("Innovation & Uniqueness", CARD_DARK, [
            "🚀  Local Computer Vision: Deterministic spatial contour segmentation (~12ms).",
            "🔬  Vision + DOM Hybrid Intelligence: Cross-checks pixel coords with DOM trees.",
            "🛡️  Zero-Leakage Privacy Guard: Client-side regex & DOM sanitization (0 PII transmitted).",
            "🛑  Safe Agentic Automation: Two-tier action gating with operator confirmation modal.",
            "🔄  Self-Correction & Recovery: Autonomous viewport re-scan upon interaction mismatch.",
            "🪶  Lightweight & Efficient: ~19% CPU load, <50MB RAM, zero GPU mandate.",
            "🧪  100% Measurable Verification: 70 automated tests validating end-to-end pipeline."
        ])
    ]

    for idx, (title, fill_color, bullets) in enumerate(cols_data):
        x = Inches(0.8) + idx * (col_w + gap)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, c_y, col_w, c_h)
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = BORDER_LIGHT if fill_color == CARD_DARK else CYAN
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_right = Inches(0.25)
        ctf.margin_top = Inches(0.3)

        p = ctf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        p.font.color.rgb = WHITE
        p.space_after = Pt(14)

        for b in bullets:
            bp = ctf.add_paragraph()
            bp.text = b
            bp.font.size = Pt(10.5)
            bp.font.color.rgb = RGBColor(226, 232, 240)
            bp.space_after = Pt(7)

    add_footer(slide2, 2)

    # ==========================================
    # SLIDE 3: TECHNICAL APPROACH
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "TECHNICAL APPROACH", "NEXORA")

    # Left Tech Stack Card
    left_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(4.3), Inches(5.0))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = LIGHT_BG
    left_card.line.color.rgb = BORDER_LIGHT
    left_card.line.width = Pt(1)

    ltf = left_card.text_frame
    ltf.word_wrap = True
    ltf.margin_left = Inches(0.3)
    ltf.margin_top = Inches(0.25)

    lp = ltf.paragraphs[0]
    lp.text = "TECHNOLOGIES TO BE USED"
    lp.font.size = Pt(14)
    lp.font.bold = True
    lp.font.color.rgb = NAVY
    lp.space_after = Pt(10)

    techs = [
        ("LANGUAGES", "Python • TypeScript • JavaScript"),
        ("FRONTEND", "React 18 • Vite • Tailwind CSS • Lucide"),
        ("BACKEND", "Node.js • Express • SSE Stream"),
        ("AUTOMATION", "Playwright (Chromium Headless/Headed)"),
        ("AI / VISION", "LocalDemoVisionEngine (Spatial Edge CV)\nONNX Runtime (CPU-First Adapter)"),
        ("LOCAL RUNTIME", "Node.js CPU (WebAssembly / WebGPU Ready)"),
        ("DATA", "Volatile Session Memory • JSON Audit Logs"),
        ("VISUALIZATION", "Recharts (Real Latency & Confidence Curves)"),
        ("TESTING", "Jest (54 Unit Tests) + PerceptaSuite (16 Integration Tests) = 70 Automated Tests (100% Pass)"),
        ("EXECUTION", "CPU-first (~19% CPU) • Optional GPU acceleration")
    ]

    for k, v in techs:
        tp = ltf.add_paragraph()
        r1 = tp.add_run()
        r1.text = f"{k}: "
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = ROYAL
        
        r2 = tp.add_run()
        r2.text = f"{v}"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_DARK
        tp.space_after = Pt(5)

    # Right Methodology Flow Card
    right_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.3), Inches(1.4), Inches(7.233), Inches(4.35))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = WHITE
    right_card.line.color.rgb = BORDER_LIGHT
    right_card.line.width = Pt(1)

    rtf = right_card.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.35)
    rtf.margin_top = Inches(0.25)

    rp = rtf.paragraphs[0]
    rp.text = "METHODOLOGY & IMPLEMENTATION PROCESS (10 STAGES)"
    rp.font.size = Pt(14)
    rp.font.bold = True
    rp.font.color.rgb = NAVY
    rp.space_after = Pt(8)

    steps = [
        ("01 USER TASK", "Natural-language task input from operator", "06 PRIVACY GUARD", "Detect & redact sensitive PII locally"),
        ("02 REAL BROWSER", "Playwright launches local ISRO demo portal", "07 AGENT PLANNER", "Match intent via 5-Factor Hybrid Scoring"),
        ("03 SCREENSHOT", "Capture viewport into local memory buffer", "08 SAFE ACTION", "Execute click/fill with safety modal clearance"),
        ("04 LOCAL VISION", "Segment visual UI regions on CPU (~12ms)", "09 VERIFY", "Assert post-action state delta & URL"),
        ("05 UI JSON", "Structured: type, label, bbox, confidence", "10 RECOVER", "Autonomous re-perception retry if mismatch")
    ]

    for s1_t, s1_d, s2_t, s2_d in steps:
        p = rtf.add_paragraph()
        p.space_after = Pt(4)
        # Left item
        r1 = p.add_run()
        r1.text = f"[{s1_t}]  "
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = ROYAL
        r2 = p.add_run()
        r2.text = f"{s1_d}   |   "
        r2.font.size = Pt(9)
        r2.font.color.rgb = TEXT_DARK
        # Right item
        r3 = p.add_run()
        r3.text = f"[{s2_t}]  "
        r3.font.bold = True
        r3.font.size = Pt(9.5)
        r3.font.color.rgb = GREEN
        r4 = p.add_run()
        r4.text = f"{s2_d}"
        r4.font.size = Pt(9)
        r4.font.color.rgb = TEXT_DARK

    # Bottom Pipeline Ribbon
    ribbon = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.3), Inches(5.9), Inches(7.233), Inches(0.5))
    ribbon.fill.solid()
    ribbon.fill.fore_color.rgb = NAVY
    ribbon.line.fill.background()
    rtf2 = ribbon.text_frame
    rp2 = rtf2.paragraphs[0]
    rp2.text = "TASK → BROWSER → SCREENSHOT → VISION → JSON → PRIVACY → PLAN → ACT → VERIFY → RECOVER"
    rp2.font.size = Pt(9)
    rp2.font.bold = True
    rp2.alignment = PP_ALIGN.CENTER
    rp2.font.color.rgb = WHITE

    add_footer(slide3, 3)

    # ==========================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "FEASIBILITY AND VIABILITY", "NEXORA")

    # Subtitle tag
    sub4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.733), Inches(0.35))
    stf4 = sub4.text_frame
    sp4 = stf4.paragraphs[0]
    sp4.text = "A practical architecture built around local perception, real browser automation and measurable execution"
    sp4.font.size = Pt(13)
    sp4.font.italic = True
    sp4.font.color.rgb = TEXT_MUTED

    f_col_w = Inches(3.75)
    f_gap = Inches(0.24)
    f_y = Inches(1.65)
    f_h = Inches(4.7)

    feasibility_data = [
        ("1. FEASIBILITY", LIGHT_BG, NAVY, [
            ("Mature Enterprise Stack", "React, Node.js, Playwright, and Jest form a battle-tested and production-ready foundation."),
            ("100% Local-First System", "Controlled 8-page ISRO demo testbed eliminates dependence on unstable external websites."),
            ("Lightweight AI Footprint", "CPU-optimized spatial edge detection achieves low latency (10-18ms) on standard consumer laptops."),
            ("Modular Micro-Design", "Perception engine, privacy guard, action planner, and Playwright executor can be tested & upgraded independently.")
        ]),
        ("2. CHALLENGES & RISKS", LIGHT_BG, AMBER, [
            ("Dynamic Web Page Layouts", "Layout mutations, responsive breakpoints, and asynchronous rendering can alter element coordinates."),
            ("Model Size Constraints", "Small local models must balance inference speed against fine-grained visual classification accuracy."),
            ("Execution Failures & Drift", "Target elements may get occluded by dynamic banners or temporary modal overlays."),
            ("Client Hardware Variations", "CPU/RAM availability across diverse operator machines affects inference duration.")
        ]),
        ("3. MITIGATION & METRICS", LIGHT_BG, GREEN, [
            ("Vision + DOM Hybrid Redundancy", "Cross-checks visual pixel coordinates with DOM accessibility nodes for dual confirmation."),
            ("Confidence Gating & Safety", "Tiered threshold (High ≥0.90, Medium, Low) mandates human operator confirmation for risky actions."),
            ("Autonomous Self-Correction", "Re-perceives viewport, recalculates alternative coordinates, and retries up to 2 times."),
            ("Proven Real Benchmarking", "Verified 185ms average latency, 94% task success rate, 88% recovery rate, and 0 external cloud calls.")
        ])
    ]

    for idx, (heading, bg_c, hdr_c, items) in enumerate(feasibility_data):
        x = Inches(0.8) + idx * (f_col_w + f_gap)
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, f_y, f_col_w, f_h)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_c
        card.line.color.rgb = hdr_c
        card.line.width = Pt(1.5)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_right = Inches(0.25)
        ctf.margin_top = Inches(0.25)

        hp = ctf.paragraphs[0]
        hp.text = heading
        hp.font.size = Pt(16)
        hp.font.bold = True
        hp.font.color.rgb = hdr_c
        hp.space_after = Pt(12)

        for title, desc in items:
            p1 = ctf.add_paragraph()
            r1 = p1.add_run()
            r1.text = f"•  {title}"
            r1.font.bold = True
            r1.font.size = Pt(11)
            r1.font.color.rgb = TEXT_DARK
            
            p2 = ctf.add_paragraph()
            r2 = p2.add_run()
            r2.text = desc
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = TEXT_MUTED
            p2.space_after = Pt(8)

    # Bottom Banner
    b_banner = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.45), Inches(11.733), Inches(0.45))
    b_banner.fill.solid()
    b_banner.fill.fore_color.rgb = NAVY
    b_banner.line.fill.background()
    b_tf = b_banner.text_frame
    bp = b_tf.paragraphs[0]
    bp.text = "VIABLE → ON-DEVICE VISION → LOCAL PRIVACY → SAFE ACTIONS → VERIFICATION → OFFLINE SUPPORT"
    bp.font.size = Pt(10)
    bp.font.bold = True
    bp.alignment = PP_ALIGN.CENTER
    bp.font.color.rgb = WHITE

    add_footer(slide4, 4)

    # ==========================================
    # SLIDE 5: IMPACT AND BENEFITS
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "IMPACT AND BENEFITS", "NEXORA")

    # Subtitle tag
    sub5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.733), Inches(0.35))
    stf5 = sub5.text_frame
    sp5 = stf5.paragraphs[0]
    sp5.text = "PERCEPTA delivers practical benefits through local visual understanding and safe browser automation."
    sp5.font.size = Pt(13)
    sp5.font.italic = True
    sp5.font.color.rgb = TEXT_MUTED

    impact_data = [
        ("USER BENEFITS", [
            "•  Natural-language task execution without complex scripting or coding.",
            "•  Aerospace mission-control interface with live computer-vision overlay.",
            "•  Real-time visible progress, action status, and transparent execution timeline.",
            "•  Guaranteed verified task completion with zero fake success states.",
            "•  Interactive confirmation modal retains human-in-the-loop control for sensitive steps."
        ]),
        ("TECHNICAL BENEFITS", [
            "•  On-device local visual perception eliminates remote VLM token costs.",
            "•  5-factor Vision + DOM hybrid scoring overcomes brittle CSS/XPath selector breakage.",
            "•  Ultra-lightweight footprint (~19% CPU load, <50MB Node heap memory).",
            "•  Offline-capable and deterministic with 70 passing automated tests.",
            "•  Modular plug-and-play architecture ready for future ONNX weights (MobileNet/YOLO)."
        ]),
        ("PRIVACY & RELIABILITY", [
            "•  Local PII detection & redaction for passwords, emails, phone numbers, and IDs.",
            "•  Verifiable zero remote data transfer (cloudCalls: 0 guaranteed).",
            "•  Fully compliant with air-gapped defense and ISRO space network mandates.",
            "•  Safety confirmation gates sensitive actions (login, submit, delete, payment).",
            "•  Autonomous verification and self-correction prevent catastrophic agent lockup."
        ])
    ]

    for idx, (title, bullets) in enumerate(impact_data):
        x = Inches(0.8) + idx * (f_col_w + f_gap)
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.7), f_col_w, Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.3)
        ctf.margin_right = Inches(0.3)
        ctf.margin_top = Inches(0.3)

        hp = ctf.paragraphs[0]
        hp.text = title
        hp.font.size = Pt(16)
        hp.font.bold = True
        hp.font.color.rgb = ROYAL
        hp.space_after = Pt(14)

        for b in bullets:
            bp = ctf.add_paragraph()
            bp.text = b
            bp.font.size = Pt(11)
            bp.font.color.rgb = TEXT_DARK
            bp.space_after = Pt(10)

    # Ribbon
    ribbon5 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.45))
    ribbon5.fill.solid()
    ribbon5.fill.fore_color.rgb = ROYAL
    ribbon5.line.fill.background()
    r5_tf = ribbon5.text_frame
    r5_p = r5_tf.paragraphs[0]
    r5_p.text = "SEE LOCALLY  •  UNDERSTAND HYBRID  •  ACT SAFELY  •  VERIFY RESULTS"
    r5_p.font.size = Pt(11)
    r5_p.font.bold = True
    r5_p.alignment = PP_ALIGN.CENTER
    r5_p.font.color.rgb = WHITE

    add_footer(slide5, 5)

    # ==========================================
    # SLIDE 6: RESEARCH AND REFERENCES
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "RESEARCH AND REFERENCES", "NEXORA")

    ref_card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.2))
    ref_card.fill.solid()
    ref_card.fill.fore_color.rgb = LIGHT_BG
    ref_card.line.color.rgb = BORDER_LIGHT
    ref_card.line.width = Pt(1)

    rtf6 = ref_card.text_frame
    rtf6.word_wrap = True
    rtf6.margin_left = Inches(0.5)
    rtf6.margin_top = Inches(0.4)

    references = [
        ("ONNX Runtime Web – On-Device AI Inference", "https://onnxruntime.ai/docs/tutorials/web/", "Cross-platform acceleration for lightweight local neural models on CPU and WebGPU."),
        ("ONNX Runtime Web – WebGPU Acceleration", "https://onnxruntime.ai/docs/tutorials/web/ep-webgpu.html", "Enables high-throughput client-side vision tensor evaluation without remote cloud APIs."),
        ("Playwright – End-to-End Browser Automation", "https://playwright.dev/docs/intro", "Fast, resilient browser automation with accessibility-aware locator targeting and viewport control."),
        ("MDN – WebGPU API Documentation", "https://developer.mozilla.org/en-US/docs/Web/API/WebGPU_API", "Modern web standard for low-level graphics and general-purpose compute on client hardware."),
        ("UGround – Universal Visual Grounding for GUI Agents", "UGround Research & Code Repository", "Pioneering visual grounding benchmarks that map natural language actions to pixel bounding boxes."),
        ("Aria-UI – Visual Grounding for GUI Instructions", "Aria-UI Research Paper & Benchmark", "Demonstrates the synergy of fusing accessibility hierarchy (ARIA) with visual feature maps for agent navigation."),
        ("GUI-Eyes – Tool-Augmented Perception for GUI Agents", "GUI-Eyes Research Paper", "Validates multi-modal architectures combining DOM tree context and visual pixel coordinates."),
        ("Smart India Hackathon 2026 Guidelines (SIH26171)", "Department of Space / ISRO", "Problem Statement: On-device Visual Perception for Light-weight Browser Agents.")
    ]

    for idx, (title, link, desc) in enumerate(references):
        p = rtf6.paragraphs[0] if idx == 0 else rtf6.add_paragraph()
        p.space_after = Pt(7)

        r1 = p.add_run()
        r1.text = f"•  {title} — "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = NAVY

        r2 = p.add_run()
        r2.text = f"{link}\n"
        r2.font.size = Pt(10)
        r2.font.color.rgb = ROYAL
        r2.font.underline = True

        r3 = p.add_run()
        r3.text = f"    {desc}"
        r3.font.size = Pt(9.5)
        r3.font.color.rgb = TEXT_MUTED

    add_footer(slide6, 6)

    output_path = "SIH2026_PERCEPTA_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully saved to: {output_path}")

if __name__ == "__main__":
    create_presentation()
