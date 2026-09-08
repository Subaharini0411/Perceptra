import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_percepta_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Reference Color Palette (Strictly extracted from Nexora.pdf)
    NAVY_TITLE = RGBColor(16, 44, 87)       # #102C57
    BLUE_SUBTITLE = RGBColor(53, 95, 172)    # #355FAC
    RED_ACCENT = RGBColor(225, 29, 72)       # #E11D48
    TEXT_BLACK = RGBColor(30, 41, 59)        # #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B
    CARD_DARK = RGBColor(60, 64, 67)         # #3C4043 Dark Grey Card
    CARD_LIGHT_BORDER = RGBColor(203, 213, 225) # #CBD5E1
    WHITE = RGBColor(255, 255, 255)
    BOTTOM_NAVY = RGBColor(24, 76, 120)      # #184C78
    CIRCLE_BLUE = RGBColor(30, 100, 180)     # Step number circles

    def add_common_header(slide, title_text):
        # Top left "NEXORA"
        t_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(2.2), Inches(0.5))
        tf = t_box.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = "NEXORA"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = NAVY_TITLE

        # Center Title
        c_box = slide.shapes.add_textbox(Inches(3.0), Inches(0.35), Inches(7.333), Inches(0.6))
        ctf = c_box.text_frame
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0
        cp = ctf.paragraphs[0]
        cp.text = title_text
        cp.font.size = Pt(24)
        cp.font.bold = True
        cp.alignment = PP_ALIGN.CENTER
        cp.font.color.rgb = NAVY_TITLE

        # Top right SIH Logo Text Badge
        r_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.25), Inches(2.2), Inches(0.8))
        rtf = r_box.text_frame
        rtf.margin_left = rtf.margin_top = rtf.margin_right = rtf.margin_bottom = 0
        rp1 = rtf.paragraphs[0]
        rp1.text = "SMART INDIA"
        rp1.font.size = Pt(11)
        rp1.font.bold = True
        rp1.alignment = PP_ALIGN.RIGHT
        rp1.font.color.rgb = NAVY_TITLE

        rp2 = rtf.add_paragraph()
        rp2.text = "HACKATHON"
        rp2.font.size = Pt(11)
        rp2.font.bold = True
        rp2.alignment = PP_ALIGN.RIGHT
        rp2.font.color.rgb = NAVY_TITLE

        rp3 = rtf.add_paragraph()
        rp3.text = "2026"
        rp3.font.size = Pt(11)
        rp3.font.bold = True
        rp3.alignment = PP_ALIGN.RIGHT
        rp3.font.color.rgb = NAVY_TITLE

    def add_common_footer(slide, page_num):
        # Bottom decorative blue line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(7.2), Inches(13.333), Inches(0.3))
        line.fill.solid()
        line.fill.fore_color.rgb = BOTTOM_NAVY
        line.line.fill.background()
        
        ltf = line.text_frame
        ltf.margin_left = Inches(0.6)
        ltf.margin_right = Inches(0.6)
        ltf.margin_top = Inches(0.03)
        lp = ltf.paragraphs[0]
        lp.text = f"@SIH Idea submission- Template                                                                                                                                                                                {page_num}"
        lp.font.size = Pt(9)
        lp.font.color.rgb = WHITE

    # ==========================================================
    # SLIDE 1: TITLE PAGE
    # ==========================================================
    slide1 = prs.slides.add_slide(blank_layout)

    # Top SIH Header
    s1_top = slide1.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(9.5), Inches(1.2))
    s1_tf = s1_top.text_frame
    s1_tf.word_wrap = True
    p = s1_tf.paragraphs[0]
    p.text = "SMART INDIA HACKATHON 2026"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = NAVY_TITLE
    p.alignment = PP_ALIGN.CENTER

    p2 = s1_tf.add_paragraph()
    p2.text = "TITLE PAGE"
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = NAVY_TITLE
    p2.alignment = PP_ALIGN.CENTER

    # SIH logo text top right
    s1_badge = slide1.shapes.add_textbox(Inches(10.5), Inches(0.35), Inches(2.2), Inches(0.8))
    s1_btf = s1_badge.text_frame
    bp1 = s1_btf.paragraphs[0]
    bp1.text = "SMART INDIA\nHACKATHON\n2026"
    bp1.font.size = Pt(11)
    bp1.font.bold = True
    bp1.alignment = PP_ALIGN.RIGHT
    bp1.font.color.rgb = NAVY_TITLE

    # Left bullet list
    list_box = slide1.shapes.add_textbox(Inches(0.6), Inches(2.3), Inches(7.2), Inches(4.5))
    ltf = list_box.text_frame
    ltf.word_wrap = True

    items = [
        ("Problem Statement ID : ", "SIH26171", False),
        ("Problem Statement Title : ", "On-device Visual\nPerception for Light-weight Browser Agents", True),
        ("Theme: ", "Miscellaneous", True),
        ("PS Category : ", "Software", True),
        ("PROJECT : ", "PERCEPTA", True),
        ("Tagline : ", "On-Device Visual Browser Agent", True),
        ("Team Name : ", "NEXORA", True)
    ]

    for i, (prefix, val, italic_val) in enumerate(items):
        par = ltf.paragraphs[0] if i == 0 else ltf.add_paragraph()
        par.space_after = Pt(14)
        
        r1 = par.add_run()
        r1.text = f"•  {prefix}"
        r1.font.bold = True
        r1.font.size = Pt(16)
        r1.font.color.rgb = TEXT_BLACK
        
        r2 = par.add_run()
        r2.text = val
        r2.font.size = Pt(16)
        r2.font.italic = italic_val
        r2.font.color.rgb = TEXT_BLACK

    # Right side graphic card
    r_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(2.1), Inches(4.3), Inches(4.6))
    r_card.fill.solid()
    r_card.fill.fore_color.rgb = RGBColor(241, 245, 249)
    r_card.line.color.rgb = CARD_LIGHT_BORDER
    r_card.line.width = Pt(1)

    rtf = r_card.text_frame
    rtf.word_wrap = True
    rtf.margin_left = rtf.margin_right = Inches(0.4)
    rtf.margin_top = Inches(0.6)

    rp = rtf.paragraphs[0]
    rp.text = "PERCEPTA"
    rp.font.size = Pt(28)
    rp.font.bold = True
    rp.alignment = PP_ALIGN.CENTER
    rp.font.color.rgb = NAVY_TITLE

    rp_sub = rtf.add_paragraph()
    rp_sub.text = "Visual Browser Agent\nISRO SIH26171 Prototype"
    rp_sub.font.size = Pt(13)
    rp_sub.alignment = PP_ALIGN.CENTER
    rp_sub.font.color.rgb = BLUE_SUBTITLE
    rp_sub.space_after = Pt(24)

    checks = [
        "✓ On-Device Computer Vision",
        "✓ Vision + DOM Hybrid Fusion",
        "✓ Privacy Guard (Zero Cloud Calls)",
        "✓ Playwright Safe Automation",
        "✓ Self-Correction & Verification"
    ]
    for c in checks:
        cp = rtf.add_paragraph()
        cp.text = c
        cp.font.size = Pt(12)
        cp.font.color.rgb = TEXT_BLACK
        cp.space_after = Pt(8)

    # ==========================================================
    # SLIDE 2: PROPOSED SOLUTION
    # ==========================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_common_header(slide2, "PERCEPTA — VISUAL BROWSER AGENT")

    # Red "Proposed Solution" Subtitle
    red_sub = slide2.shapes.add_textbox(Inches(0.6), Inches(1.1), Inches(4.0), Inches(0.4))
    rstf = red_sub.text_frame
    rstf.margin_left = rstf.margin_top = 0
    rp = rstf.paragraphs[0]
    rp.text = "Proposed Solution"
    rp.font.size = Pt(18)
    rp.font.bold = True
    rp.font.color.rgb = RED_ACCENT

    # 3 Cards Layout
    card_w = Inches(3.85)
    gap = Inches(0.24)
    c_y = Inches(1.7)
    c_h = Inches(5.1)

    cards_p2 = [
        ("Explanation", [
            ("•", "Visually understands webpages"),
            ("•", "Understands natural-language tasks"),
            ("•", "Processes perception locally on-device"),
            ("•", "Identifies the correct UI elements"),
            ("•", "Performs actions using Playwright"),
            ("•", "Protects sensitive user data"),
            ("•", "Verifies completed browser actions"),
            ("•", "Recovers from failed actions")
        ]),
        ("Problem Address", [
            ("Fragile Selectors", "Traditional automation depends heavily on XPath, CSS selectors and fixed DOM structures."),
            ("Dynamic Webpages", "Changing layouts and responsive shifts can break hard-coded automation."),
            ("Cloud Dependency", "Remote AI processing can introduce latency and severe data privacy concerns."),
            ("Silent Failures", "Conventional bots fail without verification or self-recovery capabilities.")
        ]),
        ("Innovation &\nUniqueness", [
            ("Local Computer Vision", "Visual perception runs on-device."),
            ("Vision + DOM Hybrid", "Combines visual and browser context."),
            ("Privacy Protection", "Sensitive data detected and masked locally."),
            ("Safe Agentic Automation", "Controlled actions through Playwright."),
            ("Self-Correction", "Re-perceive and retry when required."),
            ("Lightweight Architecture", "CPU-first with zero cloud LLM requirement.")
        ])
    ]

    for idx, (title, content) in enumerate(cards_p2):
        x = Inches(0.6) + idx * (card_w + gap)
        # Outer Card with rounded border
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, c_y, card_w, c_h)
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = CARD_DARK
        card.line.width = Pt(1.5)

        # Header Pill Top Shape inside Card
        hdr_shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, c_y, card_w, Inches(1.2))
        hdr_shape.fill.solid()
        hdr_shape.fill.fore_color.rgb = CARD_DARK
        hdr_shape.line.fill.background()

        htf = hdr_shape.text_frame
        htf.vertical_anchor = MSO_ANCHOR.MIDDLE
        hp = htf.paragraphs[0]
        hp.text = title
        hp.font.size = Pt(17)
        hp.font.bold = True
        hp.alignment = PP_ALIGN.CENTER
        hp.font.color.rgb = WHITE

        # Body text inside Card
        tb = slide2.shapes.add_textbox(x + Inches(0.2), c_y + Inches(1.3), card_w - Inches(0.4), c_h - Inches(1.4))
        ttf = tb.text_frame
        ttf.word_wrap = True

        if idx == 0:
            for b_idx, (bullet, text) in enumerate(content):
                p = ttf.paragraphs[0] if b_idx == 0 else ttf.add_paragraph()
                p.text = f"{bullet}  {text}"
                p.font.size = Pt(12)
                p.font.color.rgb = TEXT_BLACK
                p.space_after = Pt(10)
        else:
            for b_idx, (head, desc) in enumerate(content):
                p = ttf.paragraphs[0] if b_idx == 0 else ttf.add_paragraph()
                r1 = p.add_run()
                r1.text = f"•  {head}\n"
                r1.font.bold = True
                r1.font.size = Pt(12)
                r1.font.color.rgb = NAVY_TITLE

                r2 = p.add_run()
                r2.text = f"    {desc}"
                r2.font.size = Pt(10.5)
                r2.font.color.rgb = TEXT_MUTED
                p.space_after = Pt(8)

    add_common_footer(slide2, 2)

    # ==========================================================
    # SLIDE 3: TECHNICAL APPROACH
    # ==========================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_common_header(slide3, "PERCEPTA TECHNICAL APPROACH")

    # Left Section: Technologies To Be Used
    l_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.2), Inches(3.9), Inches(5.4))
    l_card.fill.solid()
    l_card.fill.fore_color.rgb = WHITE
    l_card.line.color.rgb = CARD_LIGHT_BORDER
    l_card.line.width = Pt(1)

    ltf = l_card.text_frame
    ltf.margin_left = ltf.margin_right = Inches(0.25)
    ltf.margin_top = Inches(0.25)
    ltf.word_wrap = True

    lp = ltf.paragraphs[0]
    lp.text = "TECHNOLOGIES TO BE USED"
    lp.font.size = Pt(13)
    lp.font.bold = True
    lp.font.color.rgb = NAVY_TITLE
    lp.space_after = Pt(10)

    techs = [
        ("LANGUAGES", "Python • TypeScript • JavaScript"),
        ("FRONTEND", "React • Vite • Tailwind CSS"),
        ("BACKEND", "Node.js • Express"),
        ("AUTOMATION", "Playwright"),
        ("AI / VISION", "• ONNX Runtime\n• Lightweight Vision Model"),
        ("LOCAL RUNTIME", "WebAssembly • WebGPU"),
        ("DATA", "JSON • SQLite"),
        ("VISUALIZATION", "Recharts"),
        ("TESTING", "Unit • Integration • E2E")
    ]

    for k, v in techs:
        tp = ltf.add_paragraph()
        r1 = tp.add_run()
        r1.text = f"{k}   "
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = BLUE_SUBTITLE

        r2 = tp.add_run()
        r2.text = f"{v}"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_BLACK
        tp.space_after = Pt(6)

    # Bottom pill in tech stack
    hw_pill = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(3.5), Inches(0.45))
    hw_pill.fill.solid()
    hw_pill.fill.fore_color.rgb = RGBColor(241, 245, 249)
    hw_pill.line.color.rgb = CARD_LIGHT_BORDER
    hwt = hw_pill.text_frame
    hwp = hwt.paragraphs[0]
    hwp.text = "CPU-first • Optional GPU acceleration"
    hwp.font.size = Pt(10)
    hwp.alignment = PP_ALIGN.CENTER
    hwp.font.color.rgb = TEXT_BLACK

    # Right Section: Methodology & Implementation Process
    r_card_p3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.7), Inches(1.2), Inches(8.0), Inches(4.8))
    r_card_p3.fill.solid()
    r_card_p3.fill.fore_color.rgb = WHITE
    r_card_p3.line.color.rgb = CARD_LIGHT_BORDER
    r_card_p3.line.width = Pt(1)

    rtf3 = r_card_p3.text_frame
    rtf3.margin_left = rtf3.margin_right = Inches(0.3)
    rtf3.margin_top = Inches(0.2)
    rtf3.word_wrap = True

    rp = rtf3.paragraphs[0]
    rp.text = "METHODOLOGY & IMPLEMENTATION PROCESS"
    rp.font.size = Pt(13)
    rp.font.bold = True
    rp.font.color.rgb = NAVY_TITLE
    rp.alignment = PP_ALIGN.CENTER
    rp.space_after = Pt(12)

    # 10 Steps in 2 columns
    steps = [
        ("01", "USER TASK", "Natural-language task input", "06", "PRIVACY GUARD", "Detect and protect sensitive data"),
        ("02", "REAL BROWSER", "Playwright opens the local demo", "07", "AGENT PLANNER", "Match intent with detected elements"),
        ("03", "SCREENSHOT", "Capture current viewport", "08", "SAFE ACTION", "Execute through Playwright"),
        ("04", "LOCAL VISION", "Detect UI elements on-device", "09", "VERIFY", "Check the new page state"),
        ("05", "UI JSON", "Type • label • bbox • confidence", "10", "RECOVER", "Re-perceive and retry when needed")
    ]

    for n1, t1, d1, n2, t2, d2 in steps:
        p = rtf3.add_paragraph()
        p.space_after = Pt(7)

        # Left Step
        r1 = p.add_run()
        r1.text = f" ({n1})  {t1} : "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = CIRCLE_BLUE

        r2 = p.add_run()
        r2.text = f"{d1}   |   "
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_BLACK

        # Right Step
        r3 = p.add_run()
        r3.text = f"({n2})  {t2} : "
        r3.font.bold = True
        r3.font.size = Pt(10)
        r3.font.color.rgb = CIRCLE_BLUE

        r4 = p.add_run()
        r4.text = f"{d2}"
        r4.font.size = Pt(9.5)
        r4.font.color.rgb = TEXT_BLACK

    # Bottom Arrow Flow Ribbon
    flow_ribbon = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.7), Inches(6.15), Inches(8.0), Inches(0.45))
    flow_ribbon.fill.solid()
    flow_ribbon.fill.fore_color.rgb = RGBColor(241, 245, 249)
    flow_ribbon.line.color.rgb = CARD_LIGHT_BORDER
    ftf = flow_ribbon.text_frame
    fp = ftf.paragraphs[0]
    fp.text = "TASK → BROWSER → SCREENSHOT → VISION → JSON → PRIVACY → PLAN → ACT → VERIFY → RECOVER"
    fp.font.size = Pt(9)
    fp.font.bold = True
    fp.alignment = PP_ALIGN.CENTER
    fp.font.color.rgb = NAVY_TITLE

    add_common_footer(slide3, 3)

    # ==========================================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # ==========================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_common_header(slide4, "PERCEPTA FEASIBILITY AND VIABILITY")

    # Subtitle
    sub_box4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.733), Inches(0.4))
    s4_tf = sub_box4.text_frame
    sp4 = s4_tf.paragraphs[0]
    sp4.text = "A practical architecture built around local perception, real browser automation and measurable execution"
    sp4.font.size = Pt(13)
    sp4.font.italic = True
    sp4.font.color.rgb = TEXT_MUTED
    sp4.alignment = PP_ALIGN.CENTER

    f_col_w = Inches(3.85)
    f_gap = Inches(0.24)
    f_y = Inches(1.6)
    f_h = Inches(4.9)

    sections = [
        ("1. FEASIBILITY", BLUE_SUBTITLE, [
            ("Mature stack", "React, Node.js, Playwright and ONNX Runtime are well suited to the prototype."),
            ("Local-first", "Controlled demo sites reduce dependence on unstable external services."),
            ("Lightweight AI", "Compact local models target CPU-friendly inference and low latency."),
            ("Modular design", "Perception, privacy, planning and execution can be tested independently.")
        ]),
        ("2. CHALLENGES & RISKS", RGBColor(194, 65, 12), [
            ("Dynamic pages", "Layouts and visual patterns can change dynamically."),
            ("Model limits", "Small models may trade accuracy for speed."),
            ("Execution failures", "Targets may disappear or actions may fail."),
            ("Hardware limits", "CPU/RAM availability affects local inference.")
        ]),
        ("3. MITIGATION", RGBColor(16, 149, 106), [
            ("Vision + DOM Hybrid", "Cross-check visual and browser information."),
            ("Confidence & safety", "Use thresholds and confirmation for risky actions."),
            ("Self-correction", "Re-perceive, select alternative and retry with limits."),
            ("Real benchmarking", "Measure latency, P50/P95, memory and model size at runtime.")
        ])
    ]

    for idx, (title, heading_color, items) in enumerate(sections):
        x = Inches(0.6) + idx * (f_col_w + f_gap)
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, f_y, f_col_w, f_h)
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = CARD_LIGHT_BORDER
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.margin_left = ctf.margin_right = Inches(0.25)
        ctf.margin_top = Inches(0.25)
        ctf.word_wrap = True

        p = ctf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = heading_color
        p.space_after = Pt(12)

        for head, desc in items:
            item_p = ctf.add_paragraph()
            r1 = item_p.add_run()
            r1.text = f"•  {head}\n"
            r1.font.bold = True
            r1.font.size = Pt(12)
            r1.font.color.rgb = NAVY_TITLE

            r2 = item_p.add_run()
            r2.text = f"    {desc}"
            r2.font.size = Pt(10)
            r2.font.color.rgb = TEXT_MUTED
            item_p.space_after = Pt(8)

    # Bottom flow ribbon
    bot4 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.65), Inches(12.133), Inches(0.38))
    bot4.fill.solid()
    bot4.fill.fore_color.rgb = RGBColor(241, 245, 249)
    bot4.line.color.rgb = CARD_LIGHT_BORDER
    b4_tf = bot4.text_frame
    b4_p = b4_tf.paragraphs[0]
    b4_p.text = "VIABLE → LOCAL VISION → PRIVACY → SAFE ACTIONS → VERIFICATION → OFFLINE SUPPORT"
    b4_p.font.size = Pt(9.5)
    b4_p.font.bold = True
    b4_p.alignment = PP_ALIGN.CENTER
    b4_p.font.color.rgb = NAVY_TITLE

    add_common_footer(slide4, 4)

    # ==========================================================
    # SLIDE 5: IMPACT AND BENEFITS
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_common_header(slide5, "PERCEPTA IMPACT AND BENEFITS")

    # Subtitle
    sub_box5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.733), Inches(0.4))
    s5_tf = sub_box5.text_frame
    sp5 = s5_tf.paragraphs[0]
    sp5.text = "PERCEPTA delivers practical benefits through local visual understanding and safe browser automation."
    sp5.font.size = Pt(13)
    sp5.font.italic = True
    sp5.font.color.rgb = TEXT_MUTED
    sp5.alignment = PP_ALIGN.CENTER

    cards_p5 = [
        ("USER BENEFITS", [
            "Natural-language task execution",
            "Simple and clear browser interaction",
            "Visible progress and action status",
            "Verified task completion",
            "Transparent confidence scoring"
        ]),
        ("TECHNICAL BENEFITS", [
            "Local visual perception",
            "Vision + DOM hybrid approach",
            "Reduced selector dependency",
            "Lightweight and offline-capable",
            "Deterministic reproducible actions"
        ]),
        ("PRIVACY & RELIABILITY", [
            "Local PII detection and redaction",
            "Reduced remote data transfer (0 calls)",
            "Safety confirmation for risky actions",
            "Verification and self-correction",
            "Air-gap & restricted network ready"
        ])
    ]

    for idx, (title, bullets) in enumerate(cards_p5):
        x = Inches(0.6) + idx * (f_col_w + f_gap)
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.6), f_col_w, Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = CARD_LIGHT_BORDER
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.margin_left = ctf.margin_right = Inches(0.3)
        ctf.margin_top = Inches(0.3)
        ctf.word_wrap = True

        p = ctf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = BLUE_SUBTITLE
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(18)

        for b in bullets:
            bp = ctf.add_paragraph()
            bp.text = f"•  {b}"
            bp.font.size = Pt(12)
            bp.font.color.rgb = TEXT_BLACK
            bp.space_after = Pt(14)

    # Bottom blue banner
    bot5 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(6.5), Inches(9.333), Inches(0.48))
    bot5.fill.solid()
    bot5.fill.fore_color.rgb = BOTTOM_NAVY
    bot5.line.fill.background()
    b5_tf = bot5.text_frame
    b5_p = b5_tf.paragraphs[0]
    b5_p.text = "SEE LOCALLY  •  ACT SAFELY  •  VERIFY RESULTS"
    b5_p.font.size = Pt(12)
    b5_p.font.bold = True
    b5_p.alignment = PP_ALIGN.CENTER
    b5_p.font.color.rgb = WHITE

    add_common_footer(slide5, 5)

    # ==========================================================
    # SLIDE 6: RESEARCH AND REFERENCES
    # ==========================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_common_header(slide6, "RESEARCH AND REFERENCES")

    # Main references container
    ref_card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.4), Inches(12.133), Inches(5.2))
    ref_card.fill.solid()
    ref_card.fill.fore_color.rgb = WHITE
    ref_card.line.color.rgb = CARD_LIGHT_BORDER
    ref_card.line.width = Pt(1)

    rtf6 = ref_card.text_frame
    rtf6.margin_left = rtf6.margin_right = Inches(0.6)
    rtf6.margin_top = Inches(0.4)
    rtf6.word_wrap = True

    references = [
        ("ONNX Runtime Web – On-Device AI Inference", "https://onnxruntime.ai/docs/tutorials/web/"),
        ("ONNX Runtime Web – WebGPU Acceleration", "https://onnxruntime.ai/docs/tutorials/web/ep-webgpu.html"),
        ("Playwright – Browser Automation", "https://playwright.dev/docs/intro"),
        ("MDN – WebGPU API", "https://developer.mozilla.org/en-US/docs/Web/API/WebGPU_API"),
        ("UGround – Universal Visual Grounding for GUI Agents", "UGround Research & Code"),
        ("Aria-UI – Visual Grounding for GUI Instructions", "Aria-UI Research Paper"),
        ("GUI-Eyes – Tool-Augmented Perception for GUI Agents", "GUI-Eyes Research Paper")
    ]

    for idx, (title, link) in enumerate(references):
        p = rtf6.paragraphs[0] if idx == 0 else rtf6.add_paragraph()
        p.space_after = Pt(14)

        r1 = p.add_run()
        r1.text = f"•   {title}   "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = TEXT_BLACK

        r2 = p.add_run()
        r2.text = link
        r2.font.size = Pt(12)
        r2.font.color.rgb = BLUE_SUBTITLE
        r2.font.underline = True

    add_common_footer(slide6, 6)

    output_path = "SIH2026_PERCEPTA_Nexora_Preserved.pptx"
    prs.save(output_path)
    print(f"Preserved presentation created: {output_path}")

if __name__ == "__main__":
    build_percepta_presentation()
